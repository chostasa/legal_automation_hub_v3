import os
import zipfile
import html
from lxml import etree
from utils.template_engine import render_docx_placeholders
from docx import Document
from core.error_handling import handle_error
from core.audit import log_audit_event
from logger import logger
import hashlib
import datetime
import re


# Import tenant and user for isolation
from core.auth import get_tenant_id, get_user_id

NAMESPACES = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

TARGET_XML_FILES = [
    "word/document.xml",
    "word/header1.xml", "word/header2.xml", "word/header3.xml",
    "word/footer1.xml", "word/footer2.xml", "word/footer3.xml",
    "word/comments.xml",
    # NOTE: Removed "vbaProject.bin" here because macro scan now handles logic directly
]


def _hash_template_version(file_path: str) -> str:
    """
    Compute SHA256 hash for template version tracking.
    """
    try:
        with open(file_path, "rb") as f:
            return hashlib.sha256(f.read()).hexdigest()
    except Exception as e:
        handle_error(e, code="DOCX_HASH_001", raise_it=True)


def _scan_for_macros(docx_path: str):
    """
    Scan the template for suspicious macros and log them.
    """
    from core.security import mask_phi, redact_log

    try:
        if not os.path.exists(docx_path):
            raise FileNotFoundError(f"File does not exist: {docx_path}")

        with zipfile.ZipFile(docx_path, 'r') as zin:
            for item in zin.infolist():
                filename_lower = item.filename.lower()
                # Only hard-fail on actual VBA macro projects
                if "vbaproject.bin" in filename_lower:
                    log_audit_event("Macro Detected", {
                        "file": docx_path,
                        "item": item.filename
                    })
                    logger.error(redact_log(mask_phi(
                        f"⚠️ Macro detected in template: {item.filename}"
                    )))
                    raise ValueError(f"Macros detected in template: {item.filename}")
                # Warn (but don't fail) for other .bin files
                elif filename_lower.endswith(".bin"):
                    logger.warning(redact_log(mask_phi(
                        f"⚠️ Non-macro .bin file found in template: {item.filename}"
                    )))
    except Exception as e:
        handle_error(e, code="DOCX_MACRO_001", raise_it=True)


def replace_text_in_docx_all(docx_path: str, replacements: dict, save_path: str) -> str:
    """
    Replace placeholders in all major parts of a Word template,
    write to a new versioned file, and log the version.
    """
    from core.security import mask_phi, redact_log

    try:
        if not isinstance(replacements, dict):
            raise ValueError("Replacements must be provided as a dictionary.")
        if not replacements:
            raise ValueError("Replacements dictionary cannot be empty.")
        if not os.path.isfile(docx_path):
            raise FileNotFoundError(f"Input DOCX file does not exist: {docx_path}")

        validate_file_size(docx_path)
        _scan_for_macros(docx_path)

        # Ensure tenant/user isolated temp directories
        tenant_id = get_tenant_id()
        user_id = get_user_id()
        save_dir = os.path.join(os.path.dirname(save_path), tenant_id, user_id)
        os.makedirs(save_dir, exist_ok=True)
        save_path = os.path.join(save_dir, os.path.basename(save_path))

        replacements = {
            k: html.unescape(str(v)) if not isinstance(v, list)
            else [html.unescape(str(x)) for x in v]
            for k, v in replacements.items()
        }

        with zipfile.ZipFile(docx_path, 'r') as zin:
            with zipfile.ZipFile(save_path, 'w') as zout:
                for item in zin.infolist():
                    buffer = zin.read(item.filename)

                    if item.filename in TARGET_XML_FILES:
                        try:
                            xml = etree.fromstring(buffer)
                            for node in xml.xpath('//w:t', namespaces=NAMESPACES):
                                if node.text:
                                    rendered = render_docx_placeholders(node.text, replacements)
                                    node.text = html.unescape(rendered)
                            buffer = etree.tostring(xml, xml_declaration=True, encoding='utf-8')
                        except Exception as e:
                            handle_error(e, code="DOCX_PARSE_001", raise_it=True)

                    zout.writestr(item, buffer)

        # Handle bullet list placeholders inside paragraphs
        doc = Document(save_path)
        for para in doc.paragraphs:
            for key, val in replacements.items():
                placeholder = f"{{{{{key}}}}}"
                if para.text.strip() == placeholder:
                    if isinstance(val, list):
                        for bullet in val:
                            if bullet.strip():
                                new_para = para.insert_paragraph_before(bullet.strip())
                                new_para.style = "List Bullet"
                        para.text = ""
                    else:
                        para.text = para.text.replace(placeholder, str(val))

        if not os.path.isfile(save_path):
            raise FileNotFoundError(f"Output DOCX was not created: {save_path}")

        doc.save(save_path)
        version_hash = _hash_template_version(save_path)

        log_audit_event("DOCX Replace Completed", {
            "file": save_path,
            "version_hash": version_hash,
            "timestamp": datetime.datetime.utcnow().isoformat(),
            "tenant_id": tenant_id,
            "user_id": user_id
        })
        logger.info(redact_log(mask_phi(
            f"✅ DOCX replacement completed for {save_path}, version {version_hash}"
        )))
        return save_path

    except Exception as e:
        handle_error(e, code="DOCX_REPLACE_001", raise_it=True)


def sanitize_filename(filename: str) -> str:
    """
    Remove invalid filename characters and return a clean name.
    """
    filename = re.sub(r'[<>:"/\\|?*]', '', filename)
    return filename.strip() or "untitled"


def validate_file_size(file_path: str, max_size_mb: int = 10) -> None:
    """
    Validate that a file is not larger than max_size_mb.
    """
    size_mb = os.path.getsize(file_path) / (1024 * 1024)
    if size_mb > max_size_mb:
        raise ValueError(
            f"File size {size_mb:.2f} MB exceeds the limit of {max_size_mb} MB."
        )


def clean_temp_dir(base_dir: str = "data/tmp") -> None:
    """
    Removes all files and folders in the temporary directory.
    """
    import shutil

    try:
        if os.path.exists(base_dir):
            shutil.rmtree(base_dir)
            # Recreate isolated directories for tenant and user
            tenant_id = get_tenant_id()
            user_id = get_user_id()
            os.makedirs(os.path.join(base_dir, tenant_id, user_id), exist_ok=True)
    except Exception as e:
        raise RuntimeError(f"Failed to clean temp directory: {e}")


from core.auth import get_tenant_id, get_user_id

def get_session_temp_dir(base_dir="data/tmp") -> str:
    from core.session_utils import get_session_id  # lazy import

    session_id = get_session_id()
    temp_dir = os.path.join(base_dir, session_id)
    os.makedirs(temp_dir, exist_ok=True)
    return temp_dir

